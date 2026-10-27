"""League division workbook: pitcher usage, team usage summary, team comparison (+ top-3 QP view).

Run from the repo root with the venv active:
    python -m scout.analysis_tools.pitching_usage --division Minors
"""

import argparse
import json
import re
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from scout.season_config import SEASON_DIR, SEASON_ID
from scout.hitting import gen_hitting
from scout.pitching.gen_pitching import (
    WALK, HBP, K_SWING, K_LOOK, HIT_OUTCOMES,
    HALF_INNING_RE, LINEUP_CHANGE_RE, PITCHER_NAMED_RE, DESC_RE, OUTCOME_KWS, SKIP_KWS,
    parse_outcome, parse_pitch_seq,
)

OUTCOME_LABELS = set(OUTCOME_KWS) | {"Catcher's Interference"}
OUTS_RE = re.compile(r"(\d)\s+Outs?")
SCORE_RE = re.compile(r"([A-Z0-9\-]{2,5})\s+(\d+)\s+-\s+([A-Z0-9\-]{2,5})\s+(\d+)")
QUALIFY_PA = 20


def _norm(first, last):
    return f"{first} {last[:1].upper()}{last[1:]}"


# ── Parsing ────────────────────────────────────────────────────────────────

def parse_game(path):
    """Return (game_id, halves) with halves in chronological order.

    Each half: {"inning", "batting", "pas": [pa, ...]} with PAs chronological.
    """
    lines = path.read_text(encoding="utf-8").splitlines()
    game_id = path.stem
    halves, cur_half, cur_pa = [], None, None

    for raw_line in lines:
        line = raw_line.strip()
        if line.startswith("GAME:"):
            m = re.search(r"/schedule/([0-9a-f-]{36})", line)
            game_id = m.group(1) if m else game_id
            continue
        m = HALF_INNING_RE.match(line)
        if m:
            cur_half = {"inning": int(m.group(2)), "half": m.group(1),
                        "batting": m.group(3).strip(), "pas": []}
            halves.append(cur_half)
            cur_pa = None
            continue
        if cur_half is None or not line or any(line.startswith(k) for k in SKIP_KWS):
            continue

        parts = line.split("|")
        label = parts[0].strip()
        if label in OUTCOME_LABELS and len(parts) >= 2:
            # Outs may sit in any middle field (e.g. "Strikeout |  | CBS- 1 - BLJY 10 | 2 Outs").
            om = OUTS_RE.search(" | ".join(parts[1:])) if len(parts) >= 3 else None
            sm = SCORE_RE.search(line)
            cur_pa = {"label": label, "outs_after": int(om.group(1)) if om else None,
                      "score": {sm.group(1): int(sm.group(2)), sm.group(3): int(sm.group(4))} if sm else None,
                      "text": [parts[-1].strip()]}
            continue

        dm = DESC_RE.match(line)
        if dm and cur_pa is not None:
            rest = dm.group(3)
            if rest.startswith("in for") or rest.startswith("in at"):
                continue
            cur_pa["batter"] = f"{dm.group(1)} {dm.group(2)}"
            cur_pa["desc"] = line
            cur_half["pas"].append(cur_pa)
            cur_pa = None
            continue
        if cur_pa is not None:
            cur_pa["text"].append(line)

    # GC files may be reverse-chronological: fix half order, then PA order per half.
    order = [(h["inning"], 0 if h["half"] == "Top" else 1) for h in halves]
    if order and order[0] > order[-1]:
        halves.reverse()
    for h in halves:
        outs = [p["outs_after"] for p in h["pas"] if p["outs_after"] is not None]
        reverse_pas = (len(outs) >= 2 and outs[0] > outs[-1]) or (
            len(outs) < 2 and order and order[0] > order[-1])
        if reverse_pas:
            h["pas"].reverse()
    return game_id, [h for h in halves if h["pas"]]


def game_date_key(path):
    try:
        return datetime.strptime(path.name[:5], "%b%d").timetuple()[1:3], path.name
    except ValueError:
        return (99, 99), path.name


# ── Aggregation ────────────────────────────────────────────────────────────

def new_app():
    return {"pitches": 0, "strikes": 0, "swings": 0, "whiffs": 0, "bf": 0, "outs": 0,
            "innings": set(), "k": 0, "bb": 0, "hbp": 0, "h": 0,
            "bb_by_inning": defaultdict(int), "bf_by_inning": defaultdict(int), "start": False}


def analyze(scorebooks):
    files = sorted(scorebooks.glob("*.txt"), key=game_date_key)
    apps = defaultdict(new_app)            # (team, pitcher, game_id) -> appearance
    team_games = defaultdict(list)         # team -> [game_id] chronological
    batting_orders = defaultdict(list)     # team -> [[initials in order], ...] per game
    batter_pa = defaultdict(lambda: defaultdict(int))  # team -> initials -> PA
    team_innings = defaultdict(lambda: defaultdict(int))  # team -> (game, inning) -> BB

    for path in files:
        game_id, halves = parse_game(path)
        teams = {h["batting"] for h in halves}
        if len(teams) != 2:
            continue
        for t in teams:
            team_games[t].append(game_id)
            batting_orders[t].append([])

        current = {}            # pitching team -> pitcher
        pending = defaultdict(list)  # pitching team -> PAs before any pitcher known
        first_seen = {}

        for h in halves:
            batting = h["batting"]
            pitching = next(t for t in teams if t != batting)
            prev_outs = 0
            for pa in h["pas"]:
                batter_pa[batting][pa["batter"]] += 1
                batting_orders[batting][-1].append(pa["batter"])

                text = " ".join(pa["text"])
                old = current.get(pitching)
                pre_change_text = ""
                lc = LINEUP_CHANGE_RE.search(text)
                if lc:
                    # Pitches thrown before a mid-PA change belong to the previous pitcher.
                    pre_change_text = text[:lc.start()]
                else:
                    lc = LINEUP_CHANGE_RE.search(pa["desc"])
                if lc:
                    current[pitching] = _norm(lc.group(1), lc.group(2))
                named = PITCHER_NAMED_RE.search(pa["desc"])
                if named:
                    current[pitching] = _norm(named.group(1), named.group(2))

                outs = 0
                if pa["outs_after"] is not None:
                    outs = max(0, pa["outs_after"] - prev_outs)
                    prev_outs = pa["outs_after"]

                pc_all = parse_pitch_seq(text)
                pc_pre = parse_pitch_seq(pre_change_text) if pre_change_text else None
                oc = parse_outcome(pa["desc"], pa["label"])
                rec = {"pc": pc_all, "pc_pre": pc_pre, "old": old, "oc": oc, "outs": outs,
                       "inning": h["inning"]}

                pitcher = current.get(pitching)
                if pitcher is None:
                    pending[pitching].append(rec)
                    continue
                if pitching not in first_seen:
                    first_seen[pitching] = pitcher
                for p in pending.pop(pitching, []):
                    _credit(apps, team_innings, pitching, pitcher, game_id, p)
                _credit(apps, team_innings, pitching, pitcher, game_id, rec)

        for team, recs in pending.items():
            for p in recs:
                _credit(apps, team_innings, team, "UNKNOWN", game_id, p)
        for team, starter in first_seen.items():
            apps[(team, starter, game_id)]["start"] = True

    return apps, team_games, batting_orders, batter_pa, team_innings


def _credit(apps, team_innings, team, pitcher, game_id, rec):
    pc, pre, old = rec["pc"], rec["pc_pre"], rec["old"]
    if pre and old and old != pitcher:
        o = apps[(team, old, game_id)]
        o["pitches"] += pre["total_pitches"]; o["strikes"] += pre["total_strikes"]
        o["swings"] += pre["total_swings"]; o["whiffs"] += pre["swing_miss"]
        pc = {k: pc[k] - pre[k] for k in ("total_pitches", "total_strikes", "total_swings", "swing_miss")}
    a = apps[(team, pitcher, game_id)]
    a["pitches"] += pc["total_pitches"]; a["strikes"] += pc["total_strikes"]
    a["swings"] += pc["total_swings"]; a["whiffs"] += pc["swing_miss"]
    a["bf"] += 1
    a["outs"] += rec["outs"]
    a["innings"].add(rec["inning"])
    oc = rec["oc"]
    a["k"] += oc in (K_SWING, K_LOOK)
    a["bb"] += oc == WALK
    a["hbp"] += oc == HBP
    a["h"] += oc in HIT_OUTCOMES
    a["bb_by_inning"][rec["inning"]] += oc == WALK
    a["bf_by_inning"][rec["inning"]] += 1
    team_innings[team][(game_id, rec["inning"])] += oc == WALK


# ── Lineups ────────────────────────────────────────────────────────────────

def lineup_stats(orders, pa_counts, collisions):
    """Distinct batting orders + game-to-game changes among qualified (>=20 PA) kids."""
    def slots(order):
        seen, out = defaultdict(int), []
        for init in order:
            n_kids = len(collisions.get(init, [])) or 1
            seen[init] += 1
            if seen[init] <= n_kids:
                out.append(f"{init}#{seen[init]}" if n_kids > 1 else init)
        return out

    def qualified(slot):
        init = slot.split("#")[0]
        n_kids = len(collisions.get(init, [])) or 1
        return pa_counts.get(init, 0) / n_kids >= QUALIFY_PA

    def consistent(a, b):
        sa, sb = set(a), set(b)
        ca, cb = [x for x in a if x in sb], [x for x in b if x in sa]
        return ca == cb

    lineups = [[s for s in slots(o) if qualified(s)] for o in orders]
    lineups = [l for l in lineups if l]
    distinct, changes = [], 0
    for i, l in enumerate(lineups):
        if not any(consistent(l, d) for d in distinct):
            distinct.append(l)
        if i and not consistent(l, lineups[i - 1]):
            changes += 1
    n_qual = len({s for l in lineups for s in l})
    return len(distinct), changes, n_qual, len(lineups)


# ── Workbook ───────────────────────────────────────────────────────────────

HEADER_FILL = PatternFill("solid", fgColor="1A2B4A")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def write_sheet(ws, headers, rows, formats):
    ws.append(headers)
    for c in ws[1]:
        c.fill, c.font = HEADER_FILL, HEADER_FONT
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    for r in rows:
        ws.append(r)
    for idx, h in enumerate(headers, 1):
        fmt = formats.get(h)
        col = get_column_letter(idx)
        if fmt:
            for cell in ws[col][1:]:
                cell.number_format = fmt
        width = max(len(str(h)) * 0.9, *(len(str(r[idx - 1])) for r in rows)) if rows else len(h)
        ws.column_dimensions[col].width = min(max(8, width + 2), 28)
    ws.row_dimensions[1].height = 42
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions


def ip_str(outs):
    return f"{outs // 3}.{outs % 3}"


def div(a, b):
    return a / b if b else None


def build_workbook(division):
    reports = SEASON_DIR / division / "Reports"
    scorebooks = reports / "Scorebooks"
    rosters = json.loads((reports / "rosters.json").read_text(encoding="utf-8"))
    apps, team_games, orders, batter_pa, team_innings = analyze(scorebooks)

    def display(team, init):
        r = rosters.get(team, {})
        coll = r.get("_collision_map", {}).get(init)
        if coll:
            return f"{init} (shared initials: {', '.join(sorted(set(coll)))})"
        e = r.get(init)
        return e["display"] if isinstance(e, dict) else init

    # Per-pitcher totals
    per_pitcher = defaultdict(lambda: {"apps": 0, "starts": 0, "pitches": 0, "bf": 0, "outs": 0,
                                       "inn_app": 0, "lim2": 0, "ge2": 0, "gt40": 0, "gt50": 0, "max_p": 0,
                                       "k": 0, "bb": 0, "hbp": 0, "h": 0, "zero_bb_inn": 0})
    team_tot = defaultdict(lambda: defaultdict(int))
    for (team, pitcher, _g), a in apps.items():
        p = per_pitcher[(team, pitcher)]
        p["apps"] += 1; p["starts"] += a["start"]
        p["pitches"] += a["pitches"]; p["bf"] += a["bf"]; p["outs"] += a["outs"]
        p["inn_app"] += len(a["innings"])
        p["lim2"] += len(a["innings"]) >= 2
        p["ge2"] += a["outs"] >= 6; p["gt40"] += a["pitches"] > 40; p["gt50"] += a["pitches"] > 50
        p["max_p"] = max(p["max_p"], a["pitches"])
        for k in ("k", "bb", "hbp", "h"):
            p[k] += a[k]
        p["zero_bb_inn"] += sum(1 for n in a["bb_by_inning"].values() if n == 0)
        t = team_tot[team]
        for k in ("pitches", "strikes", "swings", "whiffs", "bf", "outs", "k", "bb", "hbp", "h"):
            t[k] += a[k]

    wb = Workbook()

    # Tab 1 — Pitcher Usage
    ws = wb.active
    ws.title = "Pitcher Usage"
    h1 = ["Team", "Pitcher", "Games Pitched", "Starts", "Pitches", "% Team Pitches", "BF",
          "% Team BF", "IP", "% Team IP", "Innings Appeared", "Avg Pitches / App",
          "Max Pitches (1 App)", "Apps in 2 Innings (limit)", "Apps ≥ 2.0 IP", "Apps > 40 Pitches", "Apps > 50 Pitches",
          "K", "BB", "HBP", "H", "0-BB Innings"]
    rows = []
    for (team, pitcher), p in sorted(per_pitcher.items(), key=lambda kv: (kv[0][0], -kv[1]["pitches"])):
        t = team_tot[team]
        rows.append([team, display(team, pitcher), p["apps"], p["starts"], p["pitches"],
                     div(p["pitches"], t["pitches"]), p["bf"], div(p["bf"], t["bf"]),
                     ip_str(p["outs"]), div(p["outs"], t["outs"]), p["inn_app"],
                     round(p["pitches"] / p["apps"], 1), p["max_p"], p["lim2"], p["ge2"], p["gt40"], p["gt50"],
                     p["k"], p["bb"], p["hbp"], p["h"], p["zero_bb_inn"]])
    pct = "0.0%"
    write_sheet(ws, h1, rows, {"% Team Pitches": pct, "% Team BF": pct, "% Team IP": pct})

    # Tab 2 — Team Usage Summary
    ws2 = wb.create_sheet("Team Usage Summary")
    h2 = ["Team", "Games", "Pitchers Used", "Avg Pitchers / Game", "Different Starters",
          "Pitchers ≥ 10% of Pitches", "Top Pitcher % of Pitches", "Top 3 % of Pitches",
          "Effective # of Pitchers", "Apps in 2 Innings (limit)", "Apps ≥ 2.0 IP", "Apps > 40 Pitches", "Apps > 50 Pitches",
          "Max Pitches (1 App)", "Team Pitches / Game"]
    rows2 = []
    for team in sorted(team_tot):
        mine = {k: v for k, v in per_pitcher.items() if k[0] == team and k[1] != "UNKNOWN"}
        tp = team_tot[team]["pitches"]
        shares = sorted((v["pitches"] / tp for v in mine.values()), reverse=True) if tp else []
        games = len(team_games[team])
        team_apps = [a for (t, _p, _g), a in apps.items() if t == team and _p != "UNKNOWN"]
        rows2.append([team, games, len(mine), round(len(team_apps) / games, 2) if games else None,
                      sum(1 for v in mine.values() if v["starts"]),
                      sum(1 for s in shares if s >= 0.10), shares[0] if shares else None,
                      sum(shares[:3]) if shares else None,
                      round(1 / sum(s * s for s in shares), 1) if shares else None,
                      sum(v["lim2"] for v in mine.values()),
                      sum(v["ge2"] for v in mine.values()), sum(v["gt40"] for v in mine.values()),
                      sum(v["gt50"] for v in mine.values()),
                      max((v["max_p"] for v in mine.values()), default=0),
                      round(tp / games, 1) if games else None])
    write_sheet(ws2, h2, rows2, {"Top Pitcher % of Pitches": pct, "Top 3 % of Pitches": pct})

    # Tab 3 — Team Comparison
    ws3 = wb.create_sheet("Team Comparison")
    h3 = ["Team", "Games", "PA", "AVG", "OBP", "SLG", "OPS", "C%", "SM%",
          f"Qualified Batters (≥{QUALIFY_PA} PA)", "Distinct Batting Orders", "Lineup Changes (game-to-game)",
          "S%", "Whiff%", "K/9", "K/BB", "K%", "0-BB Innings", "0-BB Inning %", "BB/9", "WHIP"]
    rows3 = []
    files = sorted(scorebooks.glob("*.txt"))
    for team in sorted(team_tot):
        coll = rosters.get(team, {}).get("_collision_map", {})
        pas = []
        for f in files:
            pas.extend(gen_hitting.parse_game_for_team(str(f), team, collision_map=coll or None))
        roster = {k: v["display"] for k, v in rosters.get(team, {}).items()
                  if isinstance(v, dict) and "display" in v}
        tt = gen_hitting.compute_team_totals(gen_hitting.compute_stats(pas, roster)) or {}
        n_distinct, n_changes, n_qual, _ = lineup_stats(orders[team], batter_pa[team], coll)
        t = team_tot[team]
        ip = t["outs"] / 3
        innings = team_innings[team]
        zero_bb = sum(1 for n in innings.values() if n == 0)
        obp, slg = tt.get("obp"), tt.get("slg")
        rows3.append([team, len(team_games[team]), tt.get("pa"), tt.get("avg"), obp, slg,
                      (obp + slg) if obp is not None and slg is not None else None,
                      tt.get("c_pct"), tt.get("sm_pct"), n_qual, n_distinct, n_changes,
                      div(t["strikes"], t["pitches"]), div(t["whiffs"], t["swings"]),
                      div(t["k"] * 9, ip), div(t["k"], t["bb"]), div(t["k"], t["bf"]),
                      zero_bb, div(zero_bb, len(innings)), div(t["bb"] * 9, ip),
                      div(t["h"] + t["bb"], ip)])
    avg3 = "0.000"
    write_sheet(ws3, h3, rows3, {"AVG": avg3, "OBP": avg3, "SLG": avg3, "OPS": avg3,
                                 "C%": pct, "SM%": pct, "S%": pct, "Whiff%": pct, "K%": pct,
                                 "0-BB Inning %": pct, "K/9": "0.0", "K/BB": "0.00",
                                 "BB/9": "0.0", "WHIP": "0.00"})

    # Tab 4 — Team Comparison QP (pitching = each team's top 3 quality pitchers only)
    from scout.analysis_tools import games as td  # local: games imports this module
    from scout.analysis_tools.tourney_odds import build_ratings
    ratings = build_ratings(td.load_games(scorebooks, td.postseason(division)), apps)
    ws_qp = wb.create_sheet("Team Comparison QP")
    h_qp = h3[:12] + ["Top 3 Quality Pitchers", "Quality (≤ Lg Avg)", "% Team Pitches", "IP"] + h3[12:]
    rows_qp = []
    for row in rows3:
        team = row[0]
        top3 = [p for p in ratings["staffs"].get(team, []) if p["ip"] >= 3][:3]
        names = {p["name"] for p in top3}
        q = defaultdict(int)
        zero_bb = innings_pitched = 0
        for (t, p, _g), a in apps.items():
            if t != team or p not in names:
                continue
            for k in ("pitches", "strikes", "swings", "whiffs", "bf", "outs", "k", "bb", "h"):
                q[k] += a[k]
            innings_pitched += len(a["bb_by_inning"])
            zero_bb += sum(1 for n in a["bb_by_inning"].values() if n == 0)
        ip = q["outs"] / 3
        rows_qp.append(row[:12] + [
            ", ".join(display(team, p["name"]) for p in top3),
            sum(1 for p in top3 if p["D"] <= 1.0),
            div(q["pitches"], team_tot[team]["pitches"]), ip_str(q["outs"]),
            div(q["strikes"], q["pitches"]), div(q["whiffs"], q["swings"]),
            div(q["k"] * 9, ip), div(q["k"], q["bb"]), div(q["k"], q["bf"]),
            zero_bb, div(zero_bb, innings_pitched), div(q["bb"] * 9, ip),
            div(q["h"] + q["bb"], ip)])
    write_sheet(ws_qp, h_qp, rows_qp, {"AVG": avg3, "OBP": avg3, "SLG": avg3, "OPS": avg3,
                                       "C%": pct, "SM%": pct, "S%": pct, "Whiff%": pct, "K%": pct,
                                       "% Team Pitches": pct, "0-BB Inning %": pct, "K/9": "0.0",
                                       "K/BB": "0.00", "BB/9": "0.0", "WHIP": "0.00"})
    ws_qp.column_dimensions[get_column_letter(13)].width = 48

    from scout.season_config import load_season_yaml
    _sid, cfg = load_season_yaml()
    draft_rel = cfg["divisions"][division].get("csv")
    if draft_rel and (SEASON_DIR / draft_rel).exists():
        from scout.analysis_tools import draft_roi
        draft_roi.add_sheets(wb, division)

    # Notes
    ws4 = wb.create_sheet("Notes")
    rerun = [
        "HOW TO RE-RUN (after new games) — open Terminal and run:",
        "    cd ~/Documents/VSCode/Personal/Scout",
        "    source venv/bin/activate",
        f"    python -m scout.scraping.scrape_gc_playbyplay --division {division}",
        f"    python -m scout.scraping.scrape_gc_boxscores --division {division}",
        f"    python -m scout.analysis_tools.pitching_usage --division {division}",
        "The first two commands pull new games and update player names; the last rebuilds this workbook in place.",
        "",
    ]
    for line in rerun:
        ws4.append([line])
    ws4["A1"].font = Font(bold=True)
    for r in range(2, 7):
        ws4[f"A{r}"].font = Font(name="Courier New")
    notes = [
        f"{SEASON_ID} {division} — built {datetime.now():%Y-%m-%d %H:%M} from {len(files)} GameChanger play-by-play files.",
        "Pitches = pitches logged in the play-by-play (balls, strikes, fouls, balls in play). Pickoff throws are not counted.",
        "BF = batters faced. IP = outs recorded while pitching / 3 (shown as innings.outs, e.g. 1.2), including caught-stealing/pickoff outs.",
        "Innings Appeared = innings the pitcher threw at least one pitch in. Apps in 2 Innings (limit) = appearances spanning 2+ innings "
        "(about 1 in 3 innings end early on the run rule, so this is the true 2-inning-limit count). Apps ≥ 2.0 IP = appearances with 6+ outs recorded.",
        "Effective # of Pitchers = 1 / sum(share of team pitches²): 1.0 = one kid threw everything; higher = workload spread more evenly.",
        "0-BB Innings = innings pitched with zero walks (per pitcher on tab 1; per team on Team Comparison).",
        "S% = strikes (called + swinging + fouls + in play) / pitches. Whiff% = swinging misses / swings. WHIP = (H + BB) / IP.",
        f"Batting orders: only kids with ≥{QUALIFY_PA} season PA count; absent kids drop out instead of making a new order. "
        "Distinct = orders not matching any earlier order; Changes = games whose order differed from the previous game.",
        "Pitchers are identified by initials in GC play-by-play. Kids on the same team who share initials can't be told apart and are combined (noted in the name).",
        "UNKNOWN = pitches before GC named any pitcher in that game (rare).",
        "Team Comparison QP: hitting/lineup columns match Team Comparison; pitching columns count only each team's top 3 "
        "pitchers by the odds model's opponent-adjusted runs-allowed rating (regular season, min 3 IP). "
        "'Quality (≤ Lg Avg)' = how many of those 3 rate at or better than league average.",
        "Check: league-wide K and BB parsed here match the GameChanger box scores (K 995 vs 993, BB 1155 vs 1155).",
    ]
    for n in notes:
        ws4.append([n])
    ws4.column_dimensions["A"].width = 140

    out = reports / f"{SEASON_ID}_{division}_Pitching_Usage_and_Team_Comparison.xlsx"
    wb.save(out)
    return out, rows, rows2, rows3


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--division", default="Minors")
    args = ap.parse_args()
    out, r1, r2, r3 = build_workbook(args.division)
    print(f"Saved → {out}")
    print(f"Pitchers: {len(r1)}  Teams: {len(r2)}")
