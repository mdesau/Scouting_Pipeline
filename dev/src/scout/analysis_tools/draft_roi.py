"""Draft ROI: grade each coach's draft from the season's results.

Player value = batting runs above average (linear weights scaled to the league's run environment)
+ pitching runs saved (opponent-adjusted rating from tourney_odds.build_ratings), per team game.
ROI ADP = player's value rank among all drafted players; ROI = actual pick − ROI ADP (+ = steal).

Run from the repo root with the venv active:
    python -m scout.analysis_tools.draft_roi --division Minors
(or build all three analysis reports at once: python -m scout.analysis_tools.run_all --division Minors)
"""

import argparse
import csv
import json
import math
import re
from collections import defaultdict
from functools import lru_cache

from scout.analysis_tools import games as td
from scout.analysis_tools import pitching_usage as mur
from scout.analysis_tools.tourney_odds import build_ratings
from scout.hitting import gen_hitting
from scout.season_config import SEASON_DIR, SEASON_ID, load_season_yaml

# Run value of each batting event relative to an out (youth-ball scale; rescaled to league runs below).
WEIGHTS = {"bb": 0.55, "hbp": 0.55, "singles": 0.75, "doubles": 1.05, "triples": 1.35, "hrs": 1.65}
NICKNAMES = [{"jacob", "jake"}, {"william", "will", "liam"}, {"benjamin", "ben"}, {"nicholas", "nick"},
             {"christopher", "chris"}, {"matthew", "matt"}, {"joseph", "joey", "joe"}, {"daniel", "danny"},
             {"dheeraj", "dj"}]
GRADES = [(1.5, "A+"), (1.0, "A"), (0.5, "B+"), (0.0, "B"), (-0.5, "C+"), (-1.0, "C"), (-1.5, "D"), (-9, "F")]


def _same_first(a, b):
    a, b = a.lower(), b.lower()
    return a == b or any(a in s and b in s for s in NICKNAMES)


def load_draft(path):
    """Rows = rounds, columns = coaches; last row = team names. Snake order, column order = round 1."""
    with open(path, newline="", encoding="utf-8") as f:
        rows = [r for r in csv.reader(f)]
    coaches, teams = rows[0][1:], rows[-1][1:]
    keys = [f"{t}-{c}" for t, c in zip(teams, coaches)]
    n = len(keys)
    picks = []
    for prow in rows[1:-1]:
        rnd = int(prow[0])
        order = range(n) if rnd % 2 else range(n - 1, -1, -1)
        for slot, col in enumerate(order, 1):
            name = prow[col + 1].strip() if col + 1 < len(prow) else ""
            if name:
                picks.append({"round": rnd, "slot": slot, "team": keys[col], "coach": coaches[col], "name": name})
    for i, p in enumerate(sorted(picks, key=lambda p: (p["round"], p["slot"])), 1):
        p["pick"] = i
    return picks


def roster_index(rosters):
    """{team: {key: {display, first, last_init, jersey}}}, plus a per-team key → canonical-kid map."""
    idx, canon = {}, {}
    for team, players in rosters.items():
        entries = {}
        for key, v in players.items():
            if not isinstance(v, dict) or "display" not in v or key == "TEAM":
                continue
            m = re.match(r"([A-Za-z'\-]+)\.?\s+([A-Za-z])", v["display"])
            entries[key] = {"display": v["display"], "first": m.group(1) if m else key.split()[0],
                            "last_init": m.group(2).upper() if m else key.split()[-1][0], "jersey": v.get("jersey")}
        # GC sometimes lists one kid twice (e.g. "DJ Benu" / "Dheeraj B.", same jersey) — merge by jersey.
        by_jersey, cmap = {}, {}
        for key, e in sorted(entries.items()):
            j = e["jersey"]
            cmap[key] = by_jersey.setdefault(j, key) if j else key
        idx[team], canon[team] = entries, cmap
    return idx, canon


def match_pick(pick, idx):
    parts = pick["name"].split()
    first = parts[0]
    # Multi-word surnames ("De La Mora", "la Cour"): GC may use any word's initial.
    inits = {w[0].upper() for w in parts[1:]}
    def hits(team):
        return [k for k, e in idx.get(team, {}).items() if e["last_init"] in inits and _same_first(e["first"], first)]
    own = hits(pick["team"])
    if own:
        return pick["team"], own
    for team in idx:
        h = hits(team)
        if h:
            return team, h
    return None, []


def player_values(division, games, apps, rosters, idx, canon):
    """Per (team, canonical key): batting runs, pitching runs, PA, IP-equivalent, games."""
    sb = td.scorebooks_dir(division)
    files = sorted(sb.glob("*.txt"))
    team_games = defaultdict(int)
    for g in games:
        for t in g["teams"]:
            team_games[t] += 1
    avg_games = sum(team_games.values()) / len(team_games)

    bat = {}
    tot_w = tot_pa = 0.0
    for team in idx:
        coll = rosters.get(team, {}).get("_collision_map", {})
        pas = []
        for f in files:
            pas.extend(gen_hitting.parse_game_for_team(str(f), team, collision_map=coll or None))
        for b in gen_hitting.compute_stats(pas, {}):
            w = sum(WEIGHTS[k] * b[k] for k in WEIGHTS)
            bat[(team, b["initials"])] = (w, b)
            tot_w += w; tot_pa += b["pa"]
    total_runs = sum(sum(g["runs"].values()) for g in games)
    scale = total_runs / tot_w
    lg_w = tot_w / tot_pa

    vals = defaultdict(lambda: {"bat": 0.0, "pitch": 0.0, "pa": 0, "ip": 0.0, "ops": None, "shared": False})
    for (team, key), (w, b) in bat.items():
        ck = canon.get(team, {}).get(key, key)
        v = vals[(team, ck)]
        v["bat"] += scale * (w - lg_w * b["pa"]) * avg_games / team_games[team]
        v["pa"] += b["pa"]
        if b["obp"] is not None and b["slg"] is not None:
            v["ops"] = b["obp"] + b["slg"]

    all_games = [dict(g, playoff=False) for g in games]  # value uses every game played so far
    R = build_ratings(all_games, apps)
    for team, staff in R["staffs"].items():
        for p in staff:
            keys = _pitcher_keys(team, p["name"], rosters, idx, canon)
            for ck in keys:
                v = vals[(team, ck)]
                v["pitch"] += (1 - p["D"]) * R["mu"] * p["ip"] / len(keys) * avg_games / team_games[team]
                v["ip"] += p["ip"] / len(keys)
                v["shared"] = v["shared"] or len(keys) > 1
    return vals, R


def _pitcher_keys(team, name, rosters, idx, canon):
    keys = idx.get(team, {})
    if name in keys:
        return [canon[team].get(name, name)]
    parts = name.split()
    init = f"{parts[0][0]} {parts[-1][0].upper()}"
    coll = rosters.get(team, {}).get("_collision_map", {}).get(init)
    if coll:
        return sorted({canon[team].get(k, k) for k in coll if k in keys})
    if init in keys:
        return [canon[team].get(init, init)]
    return []


@lru_cache(maxsize=None)
def analyze(division):
    _sid, cfg = load_season_yaml()
    dcfg = cfg["divisions"][division]
    reports = SEASON_DIR / dcfg["roster_json"].rsplit("/", 1)[0]
    picks = load_draft(SEASON_DIR / dcfg["csv"])
    rosters = json.loads((SEASON_DIR / dcfg["roster_json"]).read_text(encoding="utf-8"))
    idx, canon = roster_index(rosters)
    _sb, _post, games, (apps, *_) = td.division_data(division)
    vals, R = player_values(division, games, apps, rosters, idx, canon)

    matched_keys = set()
    for p in picks:
        team, keys = match_pick(p, idx)
        p["played_for"] = team
        cks = sorted({canon[team].get(k, k) for k in keys}) if team else []
        matched_keys.update((team, k) for k in cks)
        v = {"bat": 0.0, "pitch": 0.0, "pa": 0, "ip": 0.0, "ops": None, "shared": False}
        for ck in cks:
            pv = vals.get((team, ck))
            if pv:
                for k in ("bat", "pitch", "pa", "ip"):
                    v[k] += pv[k]
                v["ops"] = pv["ops"] if v["ops"] is None else v["ops"]
                v["shared"] = v["shared"] or pv["shared"]
        p.update(v)
        p["value"] = v["bat"] + v["pitch"]
        p["coach_kid"] = p["name"].split()[-1].lower() == p["coach"].split()[-1].lower()
        p["matched"] = bool(cks)

    graded = [p for p in picks if p["matched"]]
    for rank, p in enumerate(sorted(graded, key=lambda p: -p["value"]), 1):
        p["roi_adp"] = rank
        p["roi"] = p["pick"] - rank
        p["value_round"] = math.ceil(rank / len({q["team"] for q in picks}))
    by_round = defaultdict(list)
    for p in graded:
        by_round[min(p["round"], 10)].append(p["value"])
    slot_value = {r: sum(v) / len(v) for r, v in by_round.items()}
    for p in graded:
        p["over_slot"] = p["value"] - slot_value[min(p["round"], 10)]

    teams = defaultdict(lambda: {"value": 0.0, "over": 0.0, "picks": []})
    for p in graded:
        t = teams[p["team"]]
        t["value"] += p["value"]; t["over"] += p["over_slot"]; t["picks"].append(p)
    overs = [t["over"] for t in teams.values()]
    mean = sum(overs) / len(overs)
    sd = (sum((o - mean) ** 2 for o in overs) / len(overs)) ** 0.5 or 1.0
    for t in teams.values():
        z = (t["over"] - mean) / sd
        t["z"] = z
        t["grade"] = next(g for cut, g in GRADES if z >= cut)
        t["best"] = max(t["picks"], key=lambda p: p["roi"])
        t["worst"] = min(t["picks"], key=lambda p: p["roi"])

    undrafted = sorted({(team, e["display"]) for team, es in idx.items() for k, e in es.items()
                        if (team, canon[team].get(k, k)) not in matched_keys and canon[team].get(k, k) == k})
    late = _late_additions(rosters, idx, canon, matched_keys, vals, graded, slot_value,
                           len({q["team"] for q in picks}))
    return {"picks": picks, "teams": dict(teams), "slot_value": slot_value, "undrafted": undrafted, "late": late,
            "unmatched": [p for p in picks if not p["matched"]], "reports": reports, "R": R}


# ── Output ─────────────────────────────────────────────────────────────────

LATE_MIN_GAMES = 3  # fewer = a one-game sub, not a roster addition


def _late_additions(rosters, idx, canon, matched_keys, vals, graded, slot_value, n_teams):
    """Kids not on the draft sheet who played LATE_MIN_GAMES+: valued and slotted like picks, but no ADP/ROI."""
    late = []
    for team, es in idx.items():
        for k, e in es.items():
            ck = canon[team].get(k, k)
            if ck != k or (team, ck) in matched_keys or rosters[team][k].get("games_seen", 0) < LATE_MIN_GAMES:
                continue
            v = vals.get((team, ck)) or {"bat": 0.0, "pitch": 0.0, "pa": 0, "ip": 0.0, "ops": None}
            value = v["bat"] + v["pitch"]
            slot = 1 + sum(1 for q in graded if q["value"] > value)
            vround = math.ceil(slot / n_teams)
            late.append({"name": e["display"], "team": team, "played_for": team, "late": True,
                         "coach_kid": False, "matched": True, "pick": None, "round": None, "roi": None,
                         "roi_adp": slot, "value_round": vround, "bat": v["bat"], "pitch": v["pitch"],
                         "pa": v["pa"], "ip": v["ip"], "ops": v["ops"], "value": value,
                         "over_slot": value - slot_value[min(vround, max(slot_value))]})
    return sorted(late, key=lambda p: -p["value"])


def add_sheets(wb, division):
    """Draft Board / Draft Grades / Steals & Busts tabs (called by pitching_usage.build_workbook)."""
    res = analyze(division)
    pct = "0.0"
    board = wb.create_sheet("Draft Board")
    h = ["Pick", "Round", "Drafted By", "Player", "Coach's Kid", "Played For", "PA", "OPS", "IP (halves)",
         "Bat Runs", "Pitch Runs", "Total Value", "ROI ADP", "ROI (picks)", "Value Round", "Value vs Round Avg"]
    rows = [[p["pick"], p["round"], p["team"], p["name"], "Yes" if p["coach_kid"] else "",
             p["played_for"] if p["played_for"] != p["team"] else "", p.get("pa"), p.get("ops"),
             round(p.get("ip", 0), 1), round(p["bat"], 1), round(p["pitch"], 1), round(p["value"], 1),
             p.get("roi_adp"), p.get("roi"), p.get("value_round"), round(p.get("over_slot", 0), 1)]
            for p in sorted(res["picks"], key=lambda p: p["pick"])]
    mur.write_sheet(board, h, rows, {"OPS": "0.000"})

    grades = wb.create_sheet("Draft Grades")
    hg = ["Team", "Grade", "Value Over Slot (runs)", "Total Draft Value", "Best Pick (ROI)", "Worst Pick (ROI)"]
    rows = []
    for team, t in sorted(res["teams"].items(), key=lambda kv: -kv[1]["over"]):
        rows.append([team, t["grade"], round(t["over"], 1), round(t["value"], 1),
                     f"{t['best']['name']} (Rd {t['best']['round']}, {t['best']['roi']:+d})",
                     f"{t['worst']['name']} (Rd {t['worst']['round']}, {t['worst']['roi']:+d})"])
    mur.write_sheet(grades, hg, rows, {})

    sb = wb.create_sheet("Steals & Busts")
    steals, busts, kids, late = steals_busts(res)
    hs = ["Type", "Player", "Drafted By", "Pick", "Round", "ROI ADP", "ROI (picks)", "Total Value"]
    rows = [[kind, p["name"], p["team"], p["pick"], p["round"], p["roi_adp"], p["roi"], round(p["value"], 1)]
            for kind, lst in (("STEAL", steals), ("BUST", busts), ("COACH KID", kids), ("LATE ADD", late))
            for p in lst]
    mur.write_sheet(sb, hs, rows, {})
    return res


def steals_busts(res, n=8):
    g = [p for p in res["picks"] if p["matched"]]
    steals = sorted([p for p in g if p["roi_adp"] <= 22], key=lambda p: -p["roi"])[:n]
    busts = sorted([p for p in g if p["round"] <= 4], key=lambda p: p["roi"])[:n]
    kids = sorted([p for p in g if p["coach_kid"]], key=lambda p: -p["roi"])
    return steals, busts, kids, res["late"]


# Report card text — shared by the PDF and the XLSX so both say exactly the same thing.
INTRO = ("Who drafted well? Value = runs a kid added at the plate and on the mound vs. an average "
         "player, per team game. ADP = actual pick. ROI = ADP − ROI ADP (where his production says he should "
         "have gone). Positive = steal.")
RERUN_INTRO = "To rerun after the latest games, open Terminal and run:"
SECTIONS = [
    ("Steals of the Draft", 0, "Biggest ROI among kids who produced like a top-2-round pick.", "Steals"),
    ("Busts", 1, "Rounds 1–4 picks whose production fell furthest below their draft slot.", "Busts"),
    ("Commissioner Slotting Check — Coaches' Kids", 2,
     "Coaches' kids are slotted by the commissioner. Positive ROI = slotted too low; negative = too high.",
     "Coaches' Kids"),
    ("Late Additions — Not Drafted", 3,
     f"Kids who joined after the draft and played {LATE_MIN_GAMES}+ games (one-game subs left out). ROI ADP = where "
     "their production would have slotted them; no actual pick, so no ROI. Not counted in team grades.",
     "Late Additions"),
]
METHOD_TITLE = "How the Grades Were Made"
METHOD = [
    ("Draft order.", "Snake draft; round 1 in the column order of the draft sheet, reversed every other round."),
    ("Batting value.", "Each walk/HBP/single/double/triple/HR gets a run weight; weights are scaled so the "
     "league's total equals its actual runs scored. Batting runs = weights earned minus what an average hitter "
     "would earn in the same plate appearances."),
    ("Pitching value.", "Runs saved vs. average = (1 − pitcher's opponent-adjusted runs-allowed index) × "
     "league runs per half-inning × half-innings pitched (same rating as the tournament odds board)."),
    ("Fairness.", "Value is per team game, so teams that played more games don't get extra credit. Every game "
     "played so far counts (regular season + playoffs). Traded kids count for the team that drafted them."),
    ("Grades.", "Each pick is compared with the average value of its round; team grade = total value over slot, "
     "converted to a letter by how far above/below the league average it is."),
    ("Kids sharing initials", "(play-by-play uses initials) split shared pitching value evenly; see KC-004."),
    ("Late additions.", f"Kids not on the draft sheet who played {LATE_MIN_GAMES}+ games get the same value "
     "measures and a would-be slot, but no ROI and no effect on team grades."),
]
FUN_NOTE = "For fun only — they're 9. Every kid improved this season. 🙂"


def _title(division):
    return f"WCWAA {SEASON_ID.replace('-', ' ').title()} {division} — Draft Report Card"


def _rerun_cmds(division):
    return td.rerun_commands(division)


def _grade_rows(res):
    rows = [["Team", "Grade", "Value over slot", "Best pick (ROI)", "Worst pick (ROI)"]]
    for team, t in sorted(res["teams"].items(), key=lambda kv: -kv[1]["over"]):
        rows.append([team, t["grade"], f"{t['over']:+.1f} runs",
                     f"{t['best']['name']} (Rd {t['best']['round']}, {t['best']['roi']:+d})",
                     f"{t['worst']['name']} (Rd {t['worst']['round']}, {t['worst']['roi']:+d})"])
    return rows


def _player_rows(lst):
    r = [["Player", "Drafted by", "Pick (Rd)", "ADP", "ROI ADP", "ROI", "Value", "Bat / Pitch runs"]]
    for p in lst:
        if p.get("late"):
            r.append([p["name"], p["team"], "Late add", "—", str(p["roi_adp"]), "—", f"{p['value']:+.1f}",
                      f"{p['bat']:+.1f} / {p['pitch']:+.1f}"])
            continue
        head = [p["name"] + (" *" if p["coach_kid"] else ""), p["team"], f"#{p['pick']} (Rd {p['round']})", str(p["pick"])]
        if not p["matched"]:
            r.append(head + ["—"] * 4)
            continue
        r.append(head + [str(p["roi_adp"]), f"{p['roi']:+d}", f"{p['value']:+.1f}", f"{p['bat']:+.1f} / {p['pitch']:+.1f}"])
    return r


BOARD_TITLE = "Full Draft Board"
BOARD_NOTE = "Every pick in draft order. * = coach's kid; — = not found in GameChanger."

# Excel formats: real numbers shown with a +/- sign, so filters sort them numerically.
SIGNED_INT, SIGNED_1 = "+0;-0;0", "+0.0;-0.0;0.0"
RUNS_1 = '+0.0" runs";-0.0" runs";0.0" runs"'


def _grade_rows_xlsx(res):
    rows = _grade_rows(res)
    for row, (_team, t) in zip(rows[1:], sorted(res["teams"].items(), key=lambda kv: -kv[1]["over"])):
        row[2] = round(t["over"], 1)
    return rows, {2: RUNS_1}


def _player_rows_xlsx(lst):
    """Same table as the PDF, but numbers stay numbers and Bat / Pitch runs get a column each."""
    r = [["Player", "Drafted by", "Pick (Rd)", "ADP", "ROI ADP", "ROI", "Value", "Bat runs", "Pitch runs"]]
    for p in lst:
        nums = [round(p["value"], 1), round(p["bat"], 1), round(p["pitch"], 1)]
        if p.get("late"):
            r.append([p["name"], p["team"], "Late add", None, p["roi_adp"], None] + nums)
            continue
        head = [p["name"] + (" *" if p["coach_kid"] else ""), p["team"], f"#{p['pick']} (Rd {p['round']})", p["pick"]]
        if not p["matched"]:
            r.append(head + [None] * 5)
            continue
        r.append(head + [p["roi_adp"], p["roi"]] + nums)
    return r, {5: SIGNED_INT, 6: SIGNED_1, 7: SIGNED_1, 8: SIGNED_1}


def _board(res):
    return sorted(res["picks"], key=lambda p: p["pick"])


def _unmatched_line(res):
    return ("Drafted but not found in GameChanger: "
            + ", ".join(f"{p['name']} ({p['team']})" for p in res["unmatched"])) if res["unmatched"] else None


def build_xlsx(res, division):
    """Report card as a workbook: same text and tables as the PDF, one tab per section (each filterable)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    navy = "1A2B4A"
    hdr_fill, band = PatternFill("solid", fgColor=navy), PatternFill("solid", fgColor="EEF1F6")
    wrap = Alignment(wrap_text=True, vertical="top")

    def text(ws, r, value, font=None, span=7):
        ws.cell(r, 1, value).alignment = wrap
        if font:
            ws.cell(r, 1).font = font
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=span)
        return r + 1

    def table(ws, r, rows, fmts, bold_col=None):
        for i, row in enumerate(rows):
            for c, v in enumerate(row, 1):
                cell = ws.cell(r + i, c, v)
                if i == 0:
                    cell.fill, cell.font = hdr_fill, Font(bold=True, color="FFFFFF")
                else:
                    if i % 2 == 0:
                        cell.fill = band
                    if bold_col is not None and c - 1 == bold_col:
                        cell.font = Font(bold=True)
                    if c - 1 in fmts:
                        cell.number_format = fmts[c - 1]
        ws.auto_filter.ref = f"A{r}:{get_column_letter(len(rows[0]))}{r + len(rows) - 1}"
        ws.freeze_panes = ws.cell(r + 1, 1)
        return r + len(rows) + 1

    h1, h2 = Font(bold=True, size=18, color=navy), Font(bold=True, size=13, color=navy)
    small = Font(size=9, color="808080")
    wb = Workbook()
    ws = wb.active
    ws.title = "Draft Grades"
    r = text(ws, 1, _title(division), h1)
    ws.row_dimensions[2].height = 45
    r = text(ws, r, INTRO)
    r = text(ws, r, RERUN_INTRO, small)
    for cmd in _rerun_cmds(division):
        r = text(ws, r, cmd, Font(name="Courier New", size=9, color="808080"))
    r = text(ws, r + 1, "Team Draft Grades", h2)
    table(ws, r, *_grade_rows_xlsx(res), bold_col=1)
    for col, w in zip("ABCDE", (22, 8, 16, 38, 38)):
        ws.column_dimensions[col].width = w

    lists = steals_busts(res)
    sections = [(title, lists[i], note, tab) for title, i, note, tab in SECTIONS]
    for title, rows, note, tab in sections + [(BOARD_TITLE, _board(res), BOARD_NOTE, BOARD_TITLE)]:
        ws = wb.create_sheet(tab)
        r = text(ws, 1, title, h2, span=9)
        r = text(ws, r, note, small, span=9)
        table(ws, r, *_player_rows_xlsx(rows), bold_col=5)
        for col, w in zip("ABCDEFGHI", (26, 22, 14, 8, 10, 8, 9, 10, 10)):
            ws.column_dimensions[col].width = w

    ws = wb.create_sheet(METHOD_TITLE)
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 110
    r = text(ws, 1, METHOD_TITLE, h2, span=2)
    for label, body in METHOD:
        ws.cell(r, 1, label).font = Font(bold=True)
        ws.cell(r, 1).alignment = wrap
        ws.cell(r, 2, body).alignment = wrap
        ws.row_dimensions[r].height = 15 * max(1, math.ceil(len(body) / 105))
        r += 1
    if _unmatched_line(res):
        r = text(ws, r, _unmatched_line(res), small, span=2)
    text(ws, r + 1, FUN_NOTE, small, span=2)

    out = res["reports"] / f"{SEASON_ID}_{division}_Draft_Report_Card.xlsx"
    wb.save(out)
    return out


def build_pdf(res, division):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import landscape, letter
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    navy = colors.HexColor("#1a2b4a")
    ss = getSampleStyleSheet()
    h1 = ParagraphStyle("h1", parent=ss["Title"], textColor=navy, fontSize=22)
    h2 = ParagraphStyle("h2", parent=ss["Heading2"], textColor=navy)
    body = ParagraphStyle("b", parent=ss["BodyText"], fontSize=9, leading=12)
    small = ParagraphStyle("s", parent=body, fontSize=8, leading=10, textColor=colors.grey)

    def table(rows, widths, hl_col=None):
        t = Table(rows, colWidths=[w * 72 for w in widths], repeatRows=1)
        st = [("BACKGROUND", (0, 0), (-1, 0), navy), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
              ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("FONTSIZE", (0, 0), (-1, -1), 9),
              ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#eef1f6")]),
              ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#c8ced8")), ("VALIGN", (0, 0), (-1, -1), "MIDDLE")]
        if hl_col is not None:
            st.append(("FONTNAME", (hl_col, 1), (hl_col, -1), "Helvetica-Bold"))
        t.setStyle(TableStyle(st))
        return t

    lists = steals_busts(res)
    story = [Paragraph(_title(division), h1), Paragraph(INTRO, body),
             Paragraph("<br/>".join([RERUN_INTRO] + [f"<font face='Courier'>{c}</font>" for c in _rerun_cmds(division)]),
                       small),
             Spacer(1, 8), Paragraph("Team Draft Grades", h2)]
    story += [table(_grade_rows(res), [1.7, 0.6, 1.2, 2.8, 2.8], hl_col=1), PageBreak()]
    for title, i, note, _tab in SECTIONS:
        story += [Paragraph(title, h2), Paragraph(note, small),
                  table(_player_rows(lists[i]), [1.9, 1.7, 1.0, 0.6, 0.8, 0.6, 0.7, 1.3], 5), Spacer(1, 8)]
    story += [PageBreak(), Paragraph(BOARD_TITLE, h2), Paragraph(BOARD_NOTE, small),
              table(_player_rows(_board(res)), [1.9, 1.7, 1.0, 0.6, 0.8, 0.6, 0.7, 1.3], 5)]
    story += [PageBreak(), Paragraph(METHOD_TITLE, h2)]
    for label, text in METHOD:
        story.append(Paragraph(f"<b>{label}</b> {text}", body))
    if _unmatched_line(res):
        story.append(Paragraph(_unmatched_line(res), small))
    story.append(Spacer(1, 6))
    story.append(Paragraph(FUN_NOTE, small))
    out = res["reports"] / f"{SEASON_ID}_{division}_Draft_Report_Card.pdf"
    SimpleDocTemplate(str(out), pagesize=landscape(letter), leftMargin=30, rightMargin=30, topMargin=28,
                      bottomMargin=28, title="Draft Report Card").build(story)
    return out


def main():
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--division", default="Minors")
    division = ap.parse_args().division
    res = analyze(division)
    out = build_pdf(res, division)
    card = build_xlsx(res, division)
    xlsx, *_ = mur.build_workbook(division)
    for team, t in sorted(res["teams"].items(), key=lambda kv: -kv[1]["over"]):
        print(f"  {t['grade']:3s} {team:20s} over slot {t['over']:+6.1f}  best {t['best']['name']} ({t['best']['roi']:+d})"
              f"  worst {t['worst']['name']} ({t['worst']['roi']:+d})")
    print("unmatched:", [p["name"] for p in res["unmatched"]])
    print("undrafted:", res["undrafted"])
    print(f"PDF → {out}\nReport card XLSX → {card}\nWorkbook XLSX → {xlsx}")


if __name__ == "__main__":
    main()
